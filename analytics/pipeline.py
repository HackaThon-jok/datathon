"""One-report migration. Does not use clean answers to repair dirty records."""

import argparse
import csv
import hashlib
import json
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import duckdb
from openpyxl import load_workbook

FIELDS = [f"raw_col_{i}" for i in range(1, 8)] + ["source_row_id", "ingest_row_id"]
CENT = Decimal("0.01")


def number(value, integer=False):
    try:
        n = Decimal(str(value).strip().replace(",", ""))
        if not n.is_finite() or (integer and n != n.to_integral_value()):
            raise InvalidOperation
        return int(n) if integer else n.quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        return None


def extract_workbook(path, output):
    """Keep physical worksheet rows and their original row numbers."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    if sheet.max_column != 7:
        raise ValueError("MVP supports exactly seven report columns")
    with output.open("w", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(FIELDS)
        for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
            writer.writerow(list(row) + [i, i])
    wb.close()


def reference_totals(path):
    """Independent oracle: read the old report's totals, never sum transformed rows."""
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    if len(rows) < 4 or str(rows[-1][0]).strip().lower() != "total":
        raise ValueError("Expected a final Total row in independent source workbook")
    totals = [number(rows[-1][i], integer=i != 3) for i in (1, 2, 3)]
    store_totals = [number(rows[2][i], integer=i != 3) for i in (1, 2, 3)]
    if None in totals or totals != store_totals:
        raise ValueError("Legacy store and report totals are missing or disagree")
    month = datetime.strptime(str(rows[0][1]), "%B %Y").strftime("%Y-%m-01")
    return dict(
        month=month,
        store=str(rows[2][0]),
        orders=totals[0],
        quantity=totals[1],
        sales_amount=str(totals[2]),
    )


def transform(rows, batch_id, region_map=None):
    if not rows or set(rows[0]) != set(FIELDS):
        raise ValueError("Unexpected RAW schema")
    region_map = region_map or {}
    issues, staging = [], []
    # Identical copies may be removed; conflicting copies must never be guessed.
    groups = {}
    ingest_ids = set()
    for row in rows:
        sid, iid = row["source_row_id"], row["ingest_row_id"]
        if not sid.isdigit() or not iid.isdigit():
            raise ValueError("Source and ingest IDs must be positive integers")
        if int(sid) < 1 or int(iid) < 1:
            raise ValueError("Source and ingest IDs must be positive integers")
        if iid in ingest_ids:
            issues.append(
                dict(code="DUPLICATE_INGEST_ID", source_row_id=sid, severity="ERROR")
            )
        ingest_ids.add(iid)
        groups.setdefault(int(sid), []).append(row)
    retained = []
    for sid, copies in sorted(groups.items()):
        if len(copies) > 1:
            conflict = len({tuple(r[f] for f in FIELDS[:7]) for r in copies}) != 1
            issues.append(
                dict(
                    code=(
                        "CONFLICTING_DUPLICATE" if conflict else "DUPLICATE_SOURCE_ROW"
                    ),
                    source_row_id=str(sid),
                    count=len(copies) - 1,
                    severity="ERROR" if conflict else "WARNING",
                )
            )
        retained.append(min(copies, key=lambda r: int(r["ingest_row_id"])))
    if [int(r["source_row_id"]) for r in retained] != list(range(1, len(retained) + 1)):
        raise ValueError("Missing source rows: report structure cannot be trusted")
    try:
        month = datetime.strptime(retained[0]["raw_col_2"], "%B %Y").strftime(
            "%Y-%m-01"
        )
    except ValueError as exc:
        raise ValueError("Expected month header such as January 2026") from exc
    if [retained[1][f"raw_col_{i}"] for i in (2, 3, 4)] != [
        "Order",
        "Product Quantity",
        "Total Price",
    ]:
        raise ValueError("Unexpected monthly column labels")
    store = retained[2]["raw_col_1"].strip()
    if not store or store.startswith("[") or store.lower() == "total":
        raise ValueError("Expected store summary at source row 3")
    if retained[-1]["raw_col_1"].strip().lower() != "total":
        raise ValueError("Expected final report Total")
    region = region_map.get(store)
    if not region:
        issues.append(
            dict(code="REGION_UNMAPPED", source_row_id="3", severity="WARNING")
        )
    # Summary rows are excluded from aggregation but still audited.
    for row in (retained[2], retained[-1]):
        for i, field in ((2, "orders"), (3, "quantity"), (4, "sales_amount")):
            if number(row[f"raw_col_{i}"], integer=i != 4) is None:
                issues.append(
                    dict(
                        code="INVALID_OR_MISSING_NUMBER",
                        source_row_id=row["source_row_id"],
                        field=field,
                        severity="ERROR",
                    )
                )
    for row in retained[3:-1]:
        label = row["raw_col_1"].strip()
        match = re.match(r"^\[([^\]]+)\]", label)
        if not match:
            issues.append(
                dict(
                    code="UNSUPPORTED_DETAIL_ROW",
                    source_row_id=row["source_row_id"],
                    severity="ERROR",
                )
            )
        values = [number(row[f"raw_col_{i}"], integer=i != 4) for i in (2, 3, 4)]
        for field, val in zip(("orders", "quantity", "sales_amount"), values):
            if val is None:
                issues.append(
                    dict(
                        code="INVALID_OR_MISSING_NUMBER",
                        source_row_id=row["source_row_id"],
                        field=field,
                        severity="ERROR",
                    )
                )
        staging.append(
            (
                month,
                store,
                region,
                match.group(1) if match else None,
                label,
                *values,
                batch_id,
                int(row["source_row_id"]),
                int(row["ingest_row_id"]),
            )
        )
    return staging, issues


def run(source, baseline, output_root, region_map=None, publish=False):
    source, baseline, output_root = map(Path, (source, baseline, output_root))
    run_id = (
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        + "-"
        + uuid.uuid4().hex[:8]
    )
    folder = output_root / "runs" / run_id
    folder.mkdir(parents=True)
    checksum = hashlib.sha256(source.read_bytes()).hexdigest()
    batch_id = checksum[:16]
    raw_file = folder / "raw.csv"
    if source.suffix.lower() == ".xlsx":
        extract_workbook(source, raw_file)
    else:
        raw_file.write_bytes(source.read_bytes())
    with raw_file.open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    staging, issues = transform(rows, batch_id, region_map)
    expected = reference_totals(baseline)
    con = duckdb.connect(str(folder / "migration.duckdb"))
    con.execute(
        """CREATE TABLE staging_sales(month DATE, store VARCHAR, region VARCHAR,
        item_code VARCHAR, item_label VARCHAR, orders BIGINT, quantity BIGINT,
        sales_amount DECIMAL(18,2), batch_id VARCHAR, source_row_id BIGINT, ingest_row_id BIGINT)"""
    )
    con.executemany("INSERT INTO staging_sales VALUES (?,?,?,?,?,?,?,?,?,?,?)", staging)
    con.execute(
        """CREATE TABLE mart_monthly_store AS SELECT month, store, region, batch_id,
        COUNT(*) AS detail_rows, SUM(orders)::BIGINT AS orders, SUM(quantity)::BIGINT AS quantity,
        SUM(sales_amount)::DECIMAL(18,2) AS sales_amount
        FROM staging_sales GROUP BY month, store, region, batch_id"""
    )
    actual = con.execute(
        "SELECT month,store,orders,quantity,sales_amount FROM mart_monthly_store"
    ).fetchall()
    checks = []

    def check(name, ok, observed, target):
        checks.append(
            dict(
                check=name, passed=bool(ok), actual=str(observed), expected=str(target)
            )
        )

    check("one_month_store_group", len(actual) == 1, len(actual), 1)
    if len(actual) == 1:
        for field, val in zip(
            ("month", "store", "orders", "quantity", "sales_amount"), actual[0]
        ):
            check(
                "legacy_" + field,
                str(val) == str(expected[field]),
                val,
                expected[field],
            )
    check(
        "no_critical_quality_issues",
        not any(i["severity"] == "ERROR" for i in issues),
        sum(i["severity"] == "ERROR" for i in issues),
        0,
    )
    check(
        "detail_rows_preserved",
        len(staging) == len({r["source_row_id"] for r in rows}) - 4,
        len(staging),
        len({r["source_row_id"] for r in rows}) - 4,
    )
    for table in ("staging_sales", "mart_monthly_store"):
        path = str((folder / (table + ".parquet")).resolve()).replace("'", "''")
        con.execute(f"COPY {table} TO '{path}' (FORMAT PARQUET)")
    con.execute(
        "CREATE TABLE baseline(month DATE,store VARCHAR,orders BIGINT,quantity BIGINT,sales_amount DECIMAL(18,2))"
    )
    con.execute("INSERT INTO baseline VALUES (?,?,?,?,?)", list(expected.values()))
    path = str((folder / "baseline.parquet").resolve()).replace("'", "''")
    con.execute(f"COPY baseline TO '{path}' (FORMAT PARQUET)")
    columns = [
        d[0] for d in con.execute("SELECT * FROM mart_monthly_store").description
    ]
    with (folder / "mart_monthly_store.csv").open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        writer.writerows(con.fetchall())
    con.close()
    passed = all(c["passed"] for c in checks)
    report = dict(
        run_id=run_id,
        batch_id=batch_id,
        status="PASS" if passed else "FAIL",
        scope="monthly store legacy report reproduction; regional publication requires mapping",
        source_file=source.name,
        source_sha256=checksum,
        baseline_file=baseline.name,
        baseline_sha256=hashlib.sha256(baseline.read_bytes()).hexdigest(),
        raw_rows=len(rows),
        detail_rows=len(staging),
        removed_duplicates=len(rows) - len({r["source_row_id"] for r in rows}),
        checks=checks,
        issues=issues,
        region_ready=bool(staging) and all(r[2] for r in staging),
        published=False,
        output_sha256={
            f.name: hashlib.sha256(f.read_bytes()).hexdigest()
            for f in folder.glob("*.parquet")
        },
    )
    if publish and passed:
        # Candidate files are immutable; only switch pointer after successful validation.
        pointer = output_root / "published.json"
        temp = output_root / (".pointer-" + run_id + ".tmp")
        temp.write_text(
            json.dumps(dict(run_id=run_id, path=str(folder.resolve())), indent=2)
        )
        temp.replace(pointer)
        report["published"] = True
    (folder / "validation.json").write_text(json.dumps(report, indent=2))
    print(
        json.dumps(
            dict(
                status=report["status"],
                run_id=run_id,
                folder=str(folder),
                published=report["published"],
            ),
            indent=2,
        )
    )
    return report, folder


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--source", required=True)
    p.add_argument("--baseline", required=True)
    p.add_argument("--output", default="artifacts/analytics")
    p.add_argument("--region-map")
    p.add_argument("--publish", action="store_true")
    a = p.parse_args()
    mapping = json.loads(Path(a.region_map).read_text()) if a.region_map else None
    report, _ = run(a.source, a.baseline, a.output, mapping, a.publish)
    raise SystemExit(0 if report["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
